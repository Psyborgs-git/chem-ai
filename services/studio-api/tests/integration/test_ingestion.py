"""CS-0301 acceptance tests — quarantined ingestion service.

AT-0301-1  synthetic XLSX with percentages + unknown formula cache →
           values/basis preserved, formulas never executed
AT-0301-2  active-content / oversized payloads processed in
           quarantine → typed denial recorded, nothing executed
AT-0301-3  same document twice → idempotent; no duplicate records,
           parser version retained
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from chem_studio_policy.capabilities import capabilities_for_role
from sqlalchemy.orm import Session
from workers.ingestion.types import PARSER_VERSION

from studio.auth.context import ServiceContext, load_context
from studio.domain.evidence.archive import ArchiveLimits
from studio.domain.evidence.imports import ImportService
from studio.domain.evidence.service import ArtifactService
from studio.domain.evidence.vault import Vault
from studio.errors import DomainError, ErrorCode
from studio.persistence.models import (
    Artifact,
    ExtractedRecord,
    ImportBatch,
    Principal,
    PrincipalCapability,
    Workspace,
)

pytestmark = pytest.mark.integration

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx_bytes(formula: bool = True) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Recipe"
    ws["A1"] = "component"
    ws["B1"] = "amount"
    ws["A2"] = "water"
    ws["B2"] = 0.05
    ws["B2"].number_format = "0.0%"
    if formula:
        ws["C2"] = "=B2*100"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _macro_xlsx() -> bytes:
    raw = _xlsx_bytes()
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", b"MZ-fake")
    return out.getvalue()


def _principal(session: Session, ws: Workspace, kind: str, role: str, login: str) -> Principal:
    p = Principal(workspace_id=ws.id, kind=kind, login=login, display_name=login)
    session.add(p)
    session.flush()
    for cap in sorted(capabilities_for_role(role)):
        session.add(PrincipalCapability(workspace_id=ws.id, principal_id=p.id, capability=cap))
    session.flush()
    return p


@pytest.fixture()
def ctxs(session: Session) -> tuple[ServiceContext, ServiceContext]:
    ws = Workspace(slug="w", display_name="W")
    session.add(ws)
    session.flush()
    steward = _principal(session, ws, "user", "data_steward", "ds")
    agent = _principal(session, ws, "agent", "agent", "bot")
    return (
        load_context(session, ws.id, steward.id),
        load_context(session, ws.id, agent.id),
    )


@pytest.fixture()
def services(session: Session, tmp_path: Path) -> tuple[ArtifactService, ImportService]:
    vault = Vault(tmp_path)
    artifacts = ArtifactService(
        session,
        vault,
        max_bytes=64 * 1024 * 1024,
        archive_limits=ArchiveLimits(),
    )
    # isolate=False keeps the test fast; the subprocess boundary itself
    # is covered by the worker-level unit suite.
    return artifacts, ImportService(session, vault, isolate=False)


def _committed(artifacts: ArtifactService, ctx: ServiceContext, data: bytes, name: str) -> Artifact:
    artifact = artifacts.initiate(
        ctx,
        original_name=name,
        media_type=XLSX_MIME,
        declared_size=len(data),
        declared_checksum=None,
    )
    staging = artifacts.staging_path(ctx, artifact.id)
    staging.write_bytes(data)
    return artifacts.finish(ctx, artifact.id)


class TestXlsxIngestion:
    """AT-0301-1."""

    def test_percent_and_formula_records_preserve_source_truth(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _ = ctxs
        artifacts, importer = services
        artifact = _committed(artifacts, steward, _xlsx_bytes(), "recipe.xlsx")

        batch, deduped = importer.import_artifact(steward, artifact.id)
        assert not deduped
        assert batch.status == "parsed"
        assert batch.detected_type == "xlsx"
        assert batch.parser_version == PARSER_VERSION
        assert artifact.parser_version == PARSER_VERSION

        records = importer.records(steward, batch.id)
        by_cell = {r.locator["cell"]: r for r in records}
        pct = by_cell["B2"]
        assert "percent_format_ambiguous" in pct.flags
        assert pct.payload == {"number": 0.05}  # never mass fraction 5
        assert pct.original_text == "0.05"
        assert pct.status == "proposed"  # candidate, not accepted truth
        fx = by_cell["C2"]
        assert "untrusted_formula" in fx.flags
        assert fx.original_text == "=B2*100"
        assert fx.payload is None  # never evaluated

    def test_proposed_record_reviewed_by_human(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, agent = ctxs
        artifacts, importer = services
        artifact = _committed(artifacts, steward, _xlsx_bytes(), "recipe.xlsx")
        batch, _ = importer.import_artifact(steward, artifact.id)
        record = importer.records(steward, batch.id)[0]

        accepted = importer.review_record(steward, record.id, "accepted")
        assert accepted.status == "accepted"
        assert accepted.reviewed_by == steward.principal_id
        # second review is a conflict — no silent re-decision
        with pytest.raises(DomainError) as ei:
            importer.review_record(steward, record.id, "rejected")
        assert ei.value.code == ErrorCode.CONFLICT
        # agent principals cannot review
        rec2 = importer.records(steward, batch.id)[1]
        with pytest.raises(DomainError) as ei2:
            importer.review_record(agent, rec2.id, "accepted")
        assert ei2.value.code == ErrorCode.FORBIDDEN


class TestQuarantinedDenials:
    """AT-0301-2."""

    def test_active_content_records_typed_denial_no_records(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _ = ctxs
        artifacts, importer = services
        artifact = _committed(artifacts, steward, _macro_xlsx(), "evil.xlsx")

        batch, _ = importer.import_artifact(steward, artifact.id)
        assert batch.status == "quarantined"
        assert importer.records(steward, batch.id) == []
        assert any(f["code"] == "ACTIVE_CONTENT_DENIED" for f in batch.findings)

    def test_oversized_payload_denied(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _ = ctxs
        artifacts, importer = services
        importer.limits = type(importer.limits)(max_input_bytes=128)
        artifact = _committed(artifacts, steward, _xlsx_bytes(), "big.xlsx")
        batch, _ = importer.import_artifact(steward, artifact.id)
        assert batch.status == "quarantined"
        assert any(f["code"] == "INPUT_TOO_LARGE" for f in batch.findings)


class TestIdempotentReimport:
    """AT-0301-3."""

    def test_same_bytes_twice_returns_existing_batch(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _ = ctxs
        artifacts, importer = services
        artifact = _committed(artifacts, steward, _xlsx_bytes(), "recipe.xlsx")
        first, deduped1 = importer.import_artifact(steward, artifact.id)
        second, deduped2 = importer.import_artifact(steward, artifact.id)
        assert not deduped1 and deduped2
        assert first.id == second.id
        assert first.parser_version == PARSER_VERSION
        assert session.query(ImportBatch).filter_by(workspace_id=steward.workspace_id).count() == 1
        assert (
            session.query(ExtractedRecord).filter_by(workspace_id=steward.workspace_id).count()
            == first.record_count
        )

    def test_changed_bytes_same_name_new_source_revision(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _ = ctxs
        artifacts, importer = services
        a1 = _committed(artifacts, steward, _xlsx_bytes(), "recipe.xlsx")
        a2 = _committed(artifacts, steward, _xlsx_bytes(formula=False), "recipe.xlsx")
        b1, _ = importer.import_artifact(steward, a1.id)
        b2, _ = importer.import_artifact(steward, a2.id)
        assert b1.id != b2.id
        assert b1.document_group == b2.document_group
        assert b1.source_revision == 1 and b2.source_revision == 2


class TestRightsAndScope:
    def test_denied_extraction_right_refuses(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _ = ctxs
        artifacts, importer = services
        artifact = _committed(artifacts, steward, _xlsx_bytes(), "recipe.xlsx")
        artifact.rights = {**artifact.rights, "extraction": "denied"}
        session.flush()
        with pytest.raises(DomainError) as ei:
            importer.import_artifact(steward, artifact.id)
        assert ei.value.code == ErrorCode.FORBIDDEN

    def test_uncommitted_artifact_not_parseable(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _ = ctxs
        artifacts, importer = services
        artifact = artifacts.initiate(
            ctx=steward,
            original_name="recipe.xlsx",
            media_type=XLSX_MIME,
            declared_size=None,
            declared_checksum=None,
        )
        with pytest.raises(DomainError) as ei:
            importer.import_artifact(steward, artifact.id)
        assert ei.value.code == ErrorCode.CONFLICT

    def test_agent_cannot_import(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, agent = ctxs
        artifacts, importer = services
        artifact = _committed(artifacts, steward, _xlsx_bytes(), "recipe.xlsx")
        with pytest.raises(DomainError) as ei:
            importer.import_artifact(agent, artifact.id)
        assert ei.value.code in (ErrorCode.FORBIDDEN, ErrorCode.NOT_FOUND)

    def test_quality_report_rollup(
        self,
        session: Session,
        ctxs: tuple[ServiceContext, ServiceContext],
        services: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _ = ctxs
        artifacts, importer = services
        _committed(artifacts, steward, _xlsx_bytes(), "a.xlsx")
        _committed(artifacts, steward, _macro_xlsx(), "b.xlsx")
        for art in session.query(Artifact).all():
            if art.upload_state == "committed":
                importer.import_artifact(steward, art.id)
        report = importer.quality_report(steward)
        assert report["received"] == 2
        assert report["parsed"] == 1
        assert report["quarantined"] == 1
        assert report["records_proposed"] > 0
        assert report["rights_unknown_artifacts"] >= 1
