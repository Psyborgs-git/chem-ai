"""CS-0302 acceptance tests — extraction review + evidence provenance.

AT-0302-1  ambiguous extracted identities stay visible, never
           auto-accepted (claim starts proposed)
AT-0302-2  contradicting accepted claims both remain visible with
           conditions + provenance after linking
AT-0302-3  claims carry the exact source locator/original text
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from chem_studio_policy.capabilities import capabilities_for_role
from sqlalchemy.orm import Session

from studio.auth.context import ServiceContext, load_context
from studio.domain.evidence.archive import ArchiveLimits
from studio.domain.evidence.claims import ClaimService
from studio.domain.evidence.imports import ImportService
from studio.domain.evidence.service import ArtifactService
from studio.domain.evidence.vault import Vault
from studio.errors import DomainError, ErrorCode
from studio.persistence.models import (
    Principal,
    PrincipalCapability,
    Workspace,
)

pytestmark = pytest.mark.integration

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "component"
    ws["B1"] = "amount"
    ws["A2"] = "water"
    ws["B2"] = 0.05
    ws["B2"].number_format = "0.0%"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _principal(session: Session, ws: Workspace, kind: str, role: str, login: str) -> Principal:
    p = Principal(workspace_id=ws.id, kind=kind, login=login, display_name=login)
    session.add(p)
    session.flush()
    for cap in sorted(capabilities_for_role(role)):
        session.add(PrincipalCapability(workspace_id=ws.id, principal_id=p.id, capability=cap))
    session.flush()
    return p


@pytest.fixture()
def ctxs(session: Session) -> tuple[ServiceContext, ServiceContext, ServiceContext]:
    ws = Workspace(slug="w", display_name="W")
    session.add(ws)
    session.flush()
    steward = _principal(session, ws, "user", "data_steward", "ds")
    reviewer = _principal(session, ws, "user", "scientific_reviewer", "sr")
    agent = _principal(session, ws, "agent", "agent", "bot")
    return (
        load_context(session, ws.id, steward.id),
        load_context(session, ws.id, reviewer.id),
        load_context(session, ws.id, agent.id),
    )


@pytest.fixture()
def importer(session: Session, tmp_path: Path) -> tuple[ArtifactService, ImportService]:
    vault = Vault(tmp_path)
    artifacts = ArtifactService(
        session, vault, max_bytes=64 * 1024 * 1024, archive_limits=ArchiveLimits()
    )
    return artifacts, ImportService(session, vault, isolate=False)


def _record(
    artifacts: ArtifactService,
    importer: ImportService,
    ctx: ServiceContext,
) -> tuple[str, str]:
    artifact = artifacts.initiate(
        ctx,
        original_name="recipe.xlsx",
        media_type=XLSX_MIME,
        declared_size=None,
        declared_checksum=None,
    )
    artifacts.staging_path(ctx, artifact.id).write_bytes(_xlsx_bytes())
    artifacts.finish(ctx, artifact.id)
    batch, _ = importer.import_artifact(ctx, artifact.id)
    # promote the percent cell deterministically — record order isn't
    # guaranteed across backends
    record = next(r for r in importer.records(ctx, batch.id) if r.locator.get("cell") == "B2")
    return str(batch.id), str(record.id)


class TestPromotionKeepsAmbiguityVisible:
    """AT-0302-1 + AT-0302-3."""

    def test_promoted_record_becomes_proposed_claim_with_locator(
        self,
        ctxs: tuple[ServiceContext, ServiceContext, ServiceContext],
        importer: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _, _ = ctxs
        artifacts, imp = importer
        _, record_id = _record(artifacts, imp, steward)
        svc = ClaimService(imp.db)

        claim = svc.promote_record(
            steward,
            __import__("uuid").UUID(record_id),
            subject={"material": "water", "property": "amount"},
            statement={"text": "amount present"},
        )
        assert claim.status == "proposed"  # never auto-accepted
        assert claim.kind == "document_claim"
        assert claim.locator is not None and claim.locator["cell"] == "B2"
        assert claim.locator["sheet"] == "Sheet"
        assert claim.original_text == "0.05"  # AT-0302-3
        assert claim.source_record_id is not None
        # the underlying record is marked accepted at field level
        rec = imp.get_record(steward, claim.source_record_id)
        assert rec.status == "accepted"

    def test_rejected_record_cannot_be_promoted(
        self,
        ctxs: tuple[ServiceContext, ServiceContext, ServiceContext],
        importer: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, _, _ = ctxs
        artifacts, imp = importer
        _, record_id = _record(artifacts, imp, steward)
        import uuid as _uuid

        rec_uuid = _uuid.UUID(record_id)
        imp.review_record(steward, rec_uuid, "rejected")
        with pytest.raises(DomainError) as ei:
            ClaimService(imp.db).promote_record(
                steward, rec_uuid, subject={"a": 1}, statement={"b": 2}
            )
        assert ei.value.code == ErrorCode.CONFLICT


class TestContradictionsStayVisible:
    """AT-0302-2."""

    def test_contradicting_claims_both_remain_accepted_and_visible(
        self,
        ctxs: tuple[ServiceContext, ServiceContext, ServiceContext],
        importer: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, reviewer, _agent = ctxs
        svc = ClaimService(importer[1].db)
        a = svc.create_claim(
            steward,
            kind="document_claim",
            subject={"material": "resin", "property": "tg"},
            statement={"text": "tg is 120C"},
            conditions={"method": "dsc"},
            locator={"page": 3},
        )
        b = svc.create_claim(
            steward,
            kind="measured_outcome",
            subject={"material": "resin", "property": "tg"},
            statement={"text": "tg is 98C"},
            conditions={"method": "dma", "heating_rate": "2C/min"},
        )
        a = svc.review(reviewer, a.id, "accepted")
        b = svc.review(reviewer, b.id, "accepted")

        link = svc.link(reviewer, a.id, b.id, "contradicts")
        assert link.relation == "contradicts"
        # linking is idempotent
        again = svc.link(reviewer, a.id, b.id, "contradicts")
        assert again.id == link.id
        # both claims remain accepted + visible — never demoted
        assert svc.get(reviewer, a.id).status == "accepted"
        assert svc.get(reviewer, b.id).status == "accepted"
        assert svc.get(reviewer, a.id).conditions == {"method": "dsc"}
        links = svc.links_for(reviewer, a.id)
        assert any(edge.relation == "contradicts" and edge.to_claim_id == b.id for edge in links)
        # the kinds stay distinct
        kinds = {c.kind for c in svc.claims(reviewer)}
        assert {"document_claim", "measured_outcome"} <= kinds

    def test_agent_cannot_review_or_link(
        self,
        ctxs: tuple[ServiceContext, ServiceContext, ServiceContext],
        importer: tuple[ArtifactService, ImportService],
    ) -> None:
        steward, reviewer, agent = ctxs
        svc = ClaimService(importer[1].db)
        claim = svc.create_claim(
            steward,
            kind="inferred_suggestion",
            subject={"x": 1},
            statement={"y": 2},
        )
        with pytest.raises(DomainError) as ei:
            svc.review(agent, claim.id, "accepted")
        assert ei.value.code == ErrorCode.FORBIDDEN
        a = svc.review(reviewer, claim.id, "accepted")
        with pytest.raises(DomainError):
            svc.link(agent, a.id, a.id, "supports")
