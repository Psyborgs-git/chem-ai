"""CS-0301 worker-level tests — quarantine inspection + parsers.

AT-0301-1  XLSX percentages + formula cache: values/basis preserved,
           formulas never evaluated
AT-0301-2  active content + decompression abuse denied in quarantine
"""

from __future__ import annotations

import io
import zipfile

import pytest
from workers.ingestion.limits import IngestionLimits
from workers.ingestion.parsers import parse_csv, parse_json, parse_xlsx
from workers.ingestion.quarantine import detect_type, inspect
from workers.ingestion.runner import run_parse
from workers.ingestion.types import QuarantineError

LIMS = IngestionLimits()


def _xlsx_bytes(*, percent: bool = True, formula: bool = True) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Formula"
    ws["A1"] = "component"
    ws["B1"] = "amount"
    ws["A2"] = "water"
    ws["B2"] = 0.05
    if percent:
        ws["B2"].number_format = "0.0%"
    if formula:
        ws["C2"] = "=B2*100"  # written without a cached value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_with_cached_formula() -> bytes:
    """Surgery on the sheet XML so the formula cell carries a cached
    <v> — what a spreadsheet app would have saved."""
    raw = _xlsx_bytes()
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>B2*100</f>", b"<f>B2*100</f><v>5</v>")
            dst.writestr(item, data)
    return out.getvalue()


class TestXlsxPercentAndFormulaCache:
    """AT-0301-1."""

    def test_percent_cell_flagged_never_normalized_to_fraction_5(self) -> None:
        report = parse_xlsx(_xlsx_bytes(), LIMS)
        rec = next(r for r in report.records if r.locator["cell"] == "B2")
        assert rec.locator["sheet"] == "Formula"
        assert "percent_format_ambiguous" in rec.flags
        # stored value kept verbatim — '5%' never becomes mass fraction 5
        assert rec.value == {"number": 0.05}
        assert rec.original_text == "0.05"

    def test_formula_without_cache_is_untrusted_not_evaluated(self) -> None:
        report = parse_xlsx(_xlsx_bytes(), LIMS)
        rec = next(r for r in report.records if r.locator["cell"] == "C2")
        assert "untrusted_formula" in rec.flags
        assert "missing_cached_value" in rec.flags
        assert rec.value is None  # never evaluated
        assert rec.original_text == "=B2*100"  # expression kept as text

    def test_cached_formula_result_is_read_not_computed(self) -> None:
        report = parse_xlsx(_xlsx_with_cached_formula(), LIMS)
        rec = next(r for r in report.records if r.locator["cell"] == "C2")
        assert "cached_formula_result" in rec.flags
        assert "missing_cached_value" not in rec.flags
        assert rec.value == {"cached": 5}
        assert rec.original_text == "=B2*100"


class TestCsvJsonAmbiguities:
    def test_csv_percent_and_decimal_comma_flagged(self) -> None:
        data = b'name,amount\nwater,5%\nsolvent,"1,23"\n'
        report = parse_csv(data, LIMS)
        by_text = {r.original_text: r for r in report.records}
        assert "percent_literal_ambiguous" in by_text["5%"].flags
        assert by_text["5%"].locator["header"] == "amount"
        assert "decimal_separator_ambiguous" in by_text["1,23"].flags

    def test_json_leaf_locators_preserved(self) -> None:
        data = b'{"formulation": {"water": {"amount": 0.5, "unit": "kg"}}}'
        report = parse_json(data, LIMS)
        paths = {r.locator["jsonpath"] for r in report.records}
        assert "formulation.water.amount" in paths
        assert "formulation.water.unit" in paths


class TestQuarantineDenials:
    """AT-0301-2."""

    def test_active_content_refused_before_parse(self) -> None:
        raw = _xlsx_bytes()
        src = zipfile.ZipFile(io.BytesIO(raw))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as dst:
            for item in src.infolist():
                dst.writestr(item, src.read(item.filename))
            dst.writestr("xl/vbaProject.bin", b"MZ-fake-macro-store")
        data = out.getvalue()
        assert detect_type(data, "evil.xlsm") == "xlsx"
        with pytest.raises(QuarantineError) as ei:
            inspect(data, "evil.xlsm", LIMS)
        assert ei.value.code == "ACTIVE_CONTENT_DENIED"

    def test_decompression_limit_denies_zip_bomb(self) -> None:
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            dst.writestr("xl/workbook.xml", b"<x/>")
            dst.writestr("huge.bin", b"0" * (3 * 1024 * 1024))
        tiny = IngestionLimits(max_decompressed_bytes=1024)
        with pytest.raises(QuarantineError) as ei:
            inspect(out.getvalue(), "b.xlsx", tiny)
        assert ei.value.code == "DECOMPRESSION_LIMIT"

    def test_pdf_javascript_denied(self) -> None:
        data = b"%PDF-1.4\n1 0 obj<</OpenAction<</S/JavaScript/JS(app.alert)>>>>endobj"
        with pytest.raises(QuarantineError) as ei:
            inspect(data, "a.pdf", LIMS)
        assert ei.value.code == "ACTIVE_CONTENT_DENIED"

    def test_member_count_limit(self) -> None:
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as dst:
            for i in range(20):
                dst.writestr(f"m{i}.bin", b"x")
        tiny = IngestionLimits(max_zip_members=3)
        with pytest.raises(QuarantineError) as ei:
            inspect(out.getvalue(), "a.zip", tiny)
        assert ei.value.code == "TOO_MANY_MEMBERS"


class TestSubprocessBoundary:
    def test_run_parse_crosses_the_quarantine_process(self) -> None:
        report = run_parse(_xlsx_bytes(), "f.xlsx")
        assert report.detected_type == "xlsx"
        assert report.parser_name == "openpyxl"
        assert any(r.kind == "cell" for r in report.records)
        assert not report.timed_out
